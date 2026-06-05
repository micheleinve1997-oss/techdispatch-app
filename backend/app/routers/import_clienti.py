from fastapi import APIRouter, UploadFile, File, Form, HTTPException
from fastapi.responses import StreamingResponse
from typing import List, Dict, Any
from datetime import datetime, timezone
from io import BytesIO
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from bson import ObjectId

from app.database import db
from app.routers.clienti import calcola_stato, genera_codice_cliente

router = APIRouter(prefix="/api/clienti", tags=["import"])

# Campi del template con nome colonna, campo interno, obbligatorio
TEMPLATE_FIELDS = [
    ("Ragione Sociale *",      "ragione_sociale",           True),
    ("Codice Fiscale",         "codice_fiscale",            False),
    ("Partita IVA",            "partita_iva",               False),
    ("Indirizzo Sede Legale",  "sede_legale.indirizzo",     False),
    ("CAP",                    "sede_legale.cap",           False),
    ("Città",                  "sede_legale.citta",         False),
    ("Provincia",              "sede_legale.provincia",     False),
    ("Telefono",               "sede_legale.telefono",      False),
    ("Email",                  "sede_legale.email",         False),
    ("PEC",                    "sede_legale.pec",           False),
    ("Codice SDI",             "fatturazione_elettronica.codice_sdi", False),
    ("PEC Fatturazione",       "fatturazione_elettronica.pec_fe",    False),
    ("Metodo Pagamento",       "pagamento.metodo",          False),
    ("Condizioni Pagamento",   "pagamento.condizioni",      False),
    ("IBAN",                   "pagamento.iban",            False),
    ("Note",                   "note",                      False),
]

TEMPLATE_COLUMNS = [f[0] for f in TEMPLATE_FIELDS]
FIELD_MAP = {f[0]: f[1] for f in TEMPLATE_FIELDS}


@router.get("/template")
def scarica_template():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Clienti"

    header_fill = PatternFill("solid", fgColor="1E3A5F")
    required_fill = PatternFill("solid", fgColor="2563EB")
    header_font = Font(color="FFFFFF", bold=True, size=11)

    for col_idx, (col_name, _, required) in enumerate(TEMPLATE_FIELDS, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = required_fill if required else header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[cell.column_letter].width = 22

    ws.row_dimensions[1].height = 35

    # Riga esempio
    ws.cell(row=2, column=1, value="Rossi SpA")
    ws.cell(row=2, column=2, value="12345678901")
    ws.cell(row=2, column=3, value="12345678901")
    ws.cell(row=2, column=4, value="Via Roma 1")
    ws.cell(row=2, column=5, value="20100")
    ws.cell(row=2, column=6, value="Milano")
    ws.cell(row=2, column=7, value="MI")
    ws.cell(row=2, column=8, value="+39 02 1234567")
    ws.cell(row=2, column=9, value="info@rossi.it")
    ws.cell(row=2, column=13, value="Bonifico")
    ws.cell(row=2, column=14, value="30gg FM")

    # Note istruzioni
    ws_note = wb.create_sheet("Istruzioni")
    ws_note["A1"] = "ISTRUZIONI"
    ws_note["A1"].font = Font(bold=True, size=14)
    ws_note["A3"] = "• I campi con * sono obbligatori"
    ws_note["A4"] = "• Metodo Pagamento: RiBa | SDD | Bonifico | Contanti"
    ws_note["A5"] = "• Condizioni Pagamento: 30gg FM | 60gg FM | 90gg FM | Immediato"
    ws_note["A6"] = "• Provincia: 2 lettere (es. MI, LC, BG)"
    ws_note["A7"] = "• Non modificare i nomi delle colonne nella prima riga"

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=template_clienti.xlsx"}
    )


@router.post("/preview-import")
async def preview_import(file: UploadFile = File(...)):
    """Legge il file e restituisce colonne + prime righe per il mapping."""
    content = await file.read()
    try:
        if file.filename.endswith(".csv"):
            df = pd.read_csv(BytesIO(content), dtype=str, keep_default_na=False)
        else:
            df = pd.read_excel(BytesIO(content), dtype=str, keep_default_na=False)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Errore lettura file: {str(e)}")

    df = df.fillna("")
    columns = list(df.columns)

    # Suggerisce mapping automatico
    auto_mapping: Dict[str, str] = {}
    for col in columns:
        col_clean = col.strip().lower()
        for template_col, field in FIELD_MAP.items():
            if col_clean == template_col.lower().replace(" *", ""):
                auto_mapping[col] = field
                break

    preview = df.head(3).to_dict(orient="records")

    return {
        "columns": columns,
        "auto_mapping": auto_mapping,
        "template_fields": [{"label": f[0], "field": f[1], "required": f[2]} for f in TEMPLATE_FIELDS],
        "preview": preview,
        "total_rows": len(df),
    }


@router.post("/import")
async def import_clienti(file: UploadFile = File(...), mapping: str = Form(default="")):
    """Importa clienti dal file usando il mapping colonne→campi."""
    import json

    content = await file.read()
    try:
        if file.filename.endswith(".csv"):
            df = pd.read_csv(BytesIO(content), dtype=str, keep_default_na=False)
        else:
            df = pd.read_excel(BytesIO(content), dtype=str, keep_default_na=False)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Errore lettura file: {str(e)}")

    df = df.fillna("")

    # Parsing mapping JSON
    try:
        raw_map: Dict[str, str] = json.loads(mapping) if mapping else {}
        col_map = {k: v for k, v in raw_map.items() if v and v.strip()}
    except Exception:
        col_map = {}

    # Se nessun mapping, usa auto-mapping
    if not col_map:
        for col in df.columns:
            col_clean = col.strip().lower()
            for template_col, field in FIELD_MAP.items():
                if col_clean == template_col.lower().replace(" *", ""):
                    col_map[col] = field
                    break

    results = {"importati": 0, "saltati": 0, "errori": []}

    for idx, row in df.iterrows():
        try:
            data: Dict[str, Any] = {}
            for col, field in col_map.items():
                val = str(row.get(col, "")).strip()
                if not val:
                    continue
                # Imposta campo annidato
                parts = field.split(".")
                if len(parts) == 1:
                    data[field] = val
                else:
                    if parts[0] not in data:
                        data[parts[0]] = {}
                    data[parts[0]][parts[1]] = val

            if not data.get("ragione_sociale"):
                results["saltati"] += 1
                continue

            # Controllo duplicati P.IVA / CF
            piva = data.get("partita_iva", "")
            cf = data.get("codice_fiscale", "")
            if piva or cf:
                cond = []
                if piva: cond.append({"partita_iva": piva})
                if cf: cond.append({"codice_fiscale": cf})
                esistente = await db.clienti.find_one({"$or": cond, "attivo": True})
                if esistente:
                    results["saltati"] += 1
                    results["errori"].append(
                        f"Riga {idx + 2}: '{data['ragione_sociale']}' già presente come '{esistente.get('ragione_sociale')}' ({esistente.get('codice_cliente')})"
                    )
                    continue

            # Struttura completa
            doc = {
                "ragione_sociale": data.get("ragione_sociale", ""),
                "codice_fiscale": data.get("codice_fiscale"),
                "partita_iva": data.get("partita_iva"),
                "sede_legale": data.get("sede_legale"),
                "fatturazione_elettronica": data.get("fatturazione_elettronica"),
                "contatti": [],
                "pagamento": data.get("pagamento"),
                "note": data.get("note"),
                "codice_cliente": await genera_codice_cliente(),
                "attivo": True,
                "created_at": datetime.now(timezone.utc),
                "updated_at": datetime.now(timezone.utc),
            }
            doc["stato"] = calcola_stato(doc)

            await db.clienti.insert_one(doc)
            results["importati"] += 1

        except Exception as e:
            results["errori"].append(f"Riga {idx + 2}: {str(e)}")

    return results
