from fastapi import APIRouter, UploadFile, File, Form, HTTPException
from fastapi.responses import StreamingResponse
from typing import List, Dict, Any
from datetime import datetime, timezone
from io import BytesIO
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

from app.database import db
from app.routers.tecnici import genera_codice_tecnico, SPECIALIZZAZIONI_VALIDE

router = APIRouter(prefix="/api/tecnici", tags=["import-tecnici"])

TEMPLATE_FIELDS = [
    ("Nome *",               "nome",                  True),
    ("Cognome *",            "cognome",               True),
    ("Telefono",             "telefono",              False),
    ("Email",                "email",                 False),
    ("Specializzazioni",     "specializzazioni",      False),
    ("Sede Tipo",            "sede_partenza.tipo",    False),
    ("Sede Indirizzo",       "sede_partenza.indirizzo", False),
    ("Sede CAP",             "sede_partenza.cap",     False),
    ("Sede Città",           "sede_partenza.citta",   False),
    ("Sede Provincia",       "sede_partenza.provincia", False),
    ("Patente",              "patente",               False),
    ("Mezzo Proprio",        "mezzo_proprio",         False),
    ("Stato",                "stato",                 False),
    ("Note",                 "note",                  False),
]

TEMPLATE_COLUMNS = [f[0] for f in TEMPLATE_FIELDS]
FIELD_MAP = {f[0]: f[1] for f in TEMPLATE_FIELDS}

_EXTRA_ALIASES: Dict[str, str] = {
    "nome":             "nome",
    "cognome":          "cognome",
    "telefono":         "telefono",
    "tel":              "telefono",
    "phone":            "telefono",
    "email":            "email",
    "mail":             "email",
    "e-mail":           "email",
    "specializzazioni": "specializzazioni",
    "specializzazione": "specializzazioni",
    "competenze":       "specializzazioni",
    "sede tipo":        "sede_partenza.tipo",
    "tipo partenza":    "sede_partenza.tipo",
    "sede indirizzo":   "sede_partenza.indirizzo",
    "indirizzo":        "sede_partenza.indirizzo",
    "sede cap":         "sede_partenza.cap",
    "cap":              "sede_partenza.cap",
    "sede città":       "sede_partenza.citta",
    "sede citta":       "sede_partenza.citta",
    "città":            "sede_partenza.citta",
    "citta":            "sede_partenza.citta",
    "sede provincia":   "sede_partenza.provincia",
    "provincia":        "sede_partenza.provincia",
    "patente":          "patente",
    "mezzo proprio":    "mezzo_proprio",
    "mezzo_proprio":    "mezzo_proprio",
    "stato":            "stato",
    "note":             "note",
    "notes":            "note",
}


def _auto_map_col(col: str) -> str | None:
    col_clean = col.strip().lower()
    for template_col, field in FIELD_MAP.items():
        if col_clean == template_col.lower().replace(" *", ""):
            return field
    return _EXTRA_ALIASES.get(col_clean)


def _parse_bool(val: str) -> bool:
    return val.strip().lower() in ("sì", "si", "yes", "true", "1", "x")


def _parse_specializzazioni(val: str) -> List[str]:
    parts = [p.strip() for p in val.replace(";", ",").split(",") if p.strip()]
    return [p for p in parts if p in SPECIALIZZAZIONI_VALIDE]


@router.get("/template-tecnici")
def scarica_template():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Tecnici"

    header_fill = PatternFill("solid", fgColor="1E3A5F")
    required_fill = PatternFill("solid", fgColor="2563EB")
    header_font = Font(color="FFFFFF", bold=True, size=11)

    for col_idx, (col_name, _, required) in enumerate(TEMPLATE_FIELDS, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = required_fill if required else header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[cell.column_letter].width = 22

    ws.row_dimensions[1].height = 35

    # Riga esempio
    ws.cell(row=2, column=1, value="Mario")
    ws.cell(row=2, column=2, value="Rossi")
    ws.cell(row=2, column=3, value="+39 333 1234567")
    ws.cell(row=2, column=4, value="mario.rossi@example.it")
    ws.cell(row=2, column=5, value="Caldaie e riscaldamento, Impianti idrici e termoidraulica")
    ws.cell(row=2, column=6, value="casa")
    ws.cell(row=2, column=7, value="Via Manzoni 10")
    ws.cell(row=2, column=8, value="20100")
    ws.cell(row=2, column=9, value="Milano")
    ws.cell(row=2, column=10, value="MI")
    ws.cell(row=2, column=11, value="Sì")
    ws.cell(row=2, column=12, value="No")
    ws.cell(row=2, column=13, value="ATTIVO")

    ws_note = wb.create_sheet("Istruzioni")
    ws_note["A1"] = "ISTRUZIONI"
    ws_note["A1"].font = Font(bold=True, size=14)
    ws_note["A3"] = "• I campi con * sono obbligatori"
    ws_note["A4"] = "• Specializzazioni: separare con virgola. Valori validi:"
    for i, sp in enumerate(SPECIALIZZAZIONI_VALIDE, 5):
        ws_note[f"A{i}"] = f"  - {sp}"
    row = 5 + len(SPECIALIZZAZIONI_VALIDE)
    ws_note[f"A{row}"] = "• Sede Tipo: casa | ditta"
    ws_note[f"A{row+1}"] = "• Patente / Mezzo Proprio: Sì | No"
    ws_note[f"A{row+2}"] = "• Stato: ATTIVO | NON_DISPONIBILE | IN_FERIE"
    ws_note[f"A{row+3}"] = "• Non modificare i nomi delle colonne nella prima riga"

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=template_tecnici.xlsx"}
    )


@router.post("/preview-import-tecnici")
async def preview_import(file: UploadFile = File(...)):
    content = await file.read()
    try:
        if file.filename.endswith(".csv"):
            df = pd.read_csv(BytesIO(content), dtype=str, keep_default_na=False)
        else:
            df = pd.read_excel(BytesIO(content), dtype=str, keep_default_na=False)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Errore lettura file: {str(e)}")

    df = df.fillna("")
    columns = list(df.columns)

    auto_mapping: Dict[str, str] = {}
    for col in columns:
        field = _auto_map_col(col)
        if field:
            auto_mapping[col] = field

    preview = df.head(3).to_dict(orient="records")

    return {
        "columns": columns,
        "auto_mapping": auto_mapping,
        "template_fields": [{"label": f[0], "field": f[1], "required": f[2]} for f in TEMPLATE_FIELDS],
        "preview": preview,
        "total_rows": len(df),
    }


@router.post("/import-tecnici")
async def import_tecnici(file: UploadFile = File(...), mapping: str = Form(default="")):
    import json

    content = await file.read()
    try:
        if file.filename.endswith(".csv"):
            df = pd.read_csv(BytesIO(content), dtype=str, keep_default_na=False)
        else:
            df = pd.read_excel(BytesIO(content), dtype=str, keep_default_na=False)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Errore lettura file: {str(e)}")

    df = df.fillna("")

    try:
        raw_map: Dict[str, str] = json.loads(mapping) if mapping else {}
        col_map = {k: v for k, v in raw_map.items() if v and v.strip()}
    except Exception:
        col_map = {}

    if not col_map:
        for col in df.columns:
            field = _auto_map_col(col)
            if field:
                col_map[col] = field

    results: Dict[str, Any] = {"importati": 0, "saltati": 0, "errori": []}

    for idx, row in df.iterrows():
        try:
            flat: Dict[str, Any] = {}
            for col, field in col_map.items():
                val = str(row.get(col, "")).strip()
                if not val:
                    continue
                parts = field.split(".")
                if len(parts) == 1:
                    flat[field] = val
                else:
                    if parts[0] not in flat:
                        flat[parts[0]] = {}
                    flat[parts[0]][parts[1]] = val

            nome = flat.get("nome", "")
            cognome = flat.get("cognome", "")
            if not nome or not cognome:
                results["saltati"] += 1
                continue

            # Normalizza specializzazioni
            spec_raw = flat.get("specializzazioni", "")
            if isinstance(spec_raw, str):
                specializzazioni = _parse_specializzazioni(spec_raw)
            else:
                specializzazioni = []

            # Normalizza booleani
            patente_raw = flat.get("patente", "Sì")
            mezzo_raw = flat.get("mezzo_proprio", "No")
            patente = _parse_bool(patente_raw) if isinstance(patente_raw, str) else True
            mezzo_proprio = _parse_bool(mezzo_raw) if isinstance(mezzo_raw, str) else False

            # Normalizza stato
            stato_raw = flat.get("stato", "ATTIVO").upper()
            stato_valid = {"ATTIVO", "NON_DISPONIBILE", "IN_FERIE"}
            stato = stato_raw if stato_raw in stato_valid else "ATTIVO"

            # Normalizza sede partenza
            sede_raw = flat.get("sede_partenza", {})
            tipo_raw = sede_raw.get("tipo", "ditta").lower() if sede_raw else "ditta"
            tipo = "casa" if tipo_raw == "casa" else "ditta"
            sede_partenza = {
                "tipo": tipo,
                "indirizzo": sede_raw.get("indirizzo") if sede_raw else None,
                "cap": sede_raw.get("cap") if sede_raw else None,
                "citta": sede_raw.get("citta") if sede_raw else None,
                "provincia": sede_raw.get("provincia") if sede_raw else None,
                "lat": None,
                "lng": None,
            } if sede_raw else None

            doc = {
                "nome": nome,
                "cognome": cognome,
                "telefono": flat.get("telefono") or None,
                "email": flat.get("email") or None,
                "specializzazioni": specializzazioni,
                "sede_partenza": sede_partenza,
                "patente": patente,
                "mezzo_proprio": mezzo_proprio,
                "stato": stato,
                "note": flat.get("note") or None,
                "codice_tecnico": await genera_codice_tecnico(),
                "attivo": True,
                "created_at": datetime.now(timezone.utc),
                "updated_at": datetime.now(timezone.utc),
            }

            await db.tecnici.insert_one(doc)
            results["importati"] += 1

        except Exception as e:
            results["errori"].append(f"Riga {idx + 2}: {str(e)}")

    return results
