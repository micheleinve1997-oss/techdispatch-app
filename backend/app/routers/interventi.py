from datetime import datetime, timezone
from enum import Enum
from io import BytesIO
import re
from typing import Any, Dict, List, Optional

from bson import ObjectId, errors as bson_errors
from fastapi import APIRouter, File, Form, HTTPException, Query, UploadFile, status
from fastapi.responses import StreamingResponse
import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from pydantic import BaseModel, ConfigDict

from app.database import db

router = APIRouter(prefix="/api/interventi", tags=["interventi"])


class OrigineIntervento(str, Enum):
    ticket = "TICKET"
    manutenzione = "MANUTENZIONE"
    preventivo = "PREVENTIVO_ACCETTATO"
    manuale = "MANUALE"


class StatoIntervento(str, Enum):
    aperto = "APERTO"
    da_pianificare = "DA_PIANIFICARE"
    pianificato = "PIANIFICATO"
    in_corso = "IN_CORSO"
    completato = "COMPLETATO"
    annullato = "ANNULLATO"


class PrioritaIntervento(str, Enum):
    bassa = "BASSA"
    media = "MEDIA"
    alta = "ALTA"
    urgente = "URGENTE"


class InterventoCreate(BaseModel):
    model_config = ConfigDict(from_attributes=True)
    titolo: str
    descrizione: Optional[str] = None
    origine: OrigineIntervento = OrigineIntervento.manuale
    stato: StatoIntervento = StatoIntervento.aperto
    priorita: PrioritaIntervento = PrioritaIntervento.media
    durata_stimata: Optional[float] = None
    cliente_id: Optional[str] = None
    cliente_codice: Optional[str] = None
    cliente_nome: Optional[str] = None
    tecnico_id: Optional[str] = None
    tecnico_nome: Optional[str] = None
    indirizzo: Optional[str] = None
    cap: Optional[str] = None
    citta: Optional[str] = None
    provincia: Optional[str] = None
    data_richiesta: Optional[datetime] = None
    data_pianificata: Optional[datetime] = None
    lat: Optional[float] = None
    lng: Optional[float] = None
    riferimento_esterno: Optional[str] = None
    note: Optional[str] = None


class InterventoUpdate(InterventoCreate):
    titolo: Optional[str] = None


class InterventoResponse(InterventoCreate):
    id: str
    codice_intervento: str
    attivo: bool
    created_at: datetime
    updated_at: datetime


async def genera_codice_intervento() -> str:
    ultimo = await db.interventi.find_one(
        {"codice_intervento": {"$exists": True}},
        sort=[("codice_intervento", -1)]
    )
    if not ultimo:
        return "INT-0001"
    try:
        n = int(ultimo["codice_intervento"].split("-")[1])
        return f"INT-{(n + 1):04d}"
    except Exception:
        count = await db.interventi.count_documents({})
        return f"INT-{(count + 1):04d}"


TEMPLATE_FIELDS = [
    ("Titolo *", "titolo", True),
    ("Descrizione", "descrizione", False),
    ("Origine", "origine", False),
    ("Stato", "stato", False),
    ("Priorita", "priorita", False),
    ("Cliente", "cliente_nome", False),
    ("Codice Cliente", "cliente_codice", False),
    ("Tecnico", "tecnico_nome", False),
    ("Indirizzo", "indirizzo", False),
    ("CAP", "cap", False),
    ("Citta", "citta", False),
    ("Provincia", "provincia", False),
    ("Data Richiesta", "data_richiesta", False),
    ("Scadenza", "data_pianificata", False),
    ("Riferimento Esterno", "riferimento_esterno", False),
    ("Note", "note", False),
]

FIELD_MAP = {f[0]: f[1] for f in TEMPLATE_FIELDS}

_EXTRA_ALIASES: Dict[str, str] = {
    "titolo": "titolo",
    "intervento": "titolo",
    "descrizione": "descrizione",
    "descrizione intervento": "descrizione",
    "origine": "origine",
    "tipo": "origine",
    "stato": "stato",
    "priorita": "priorita",
    "priorità": "priorita",
    "cliente": "cliente_nome",
    "cliente_nome": "cliente_nome",
    "ragione sociale": "cliente_nome",
    "ragione_sociale": "cliente_nome",
    "azienda": "cliente_nome",
    "denominazione sociale": "cliente_nome",
    "codice cliente": "cliente_codice",
    "codice_cliente": "cliente_codice",
    "tecnico": "tecnico_nome",
    "tecnico_nome": "tecnico_nome",
    "indirizzo": "indirizzo",
    "via": "indirizzo",
    "sede": "indirizzo",
    "cap": "cap",
    "codice postale": "cap",
    "c.a.p.": "cap",
    "citta": "citta",
    "città": "citta",
    "comune": "citta",
    "localita": "citta",
    "località": "citta",
    "provincia": "provincia",
    "prov": "provincia",
    "pr": "provincia",
    "data richiesta": "data_richiesta",
    "data_richiesta": "data_richiesta",
    "scadenza": "data_pianificata",
    "data scadenza": "data_pianificata",
    "data pianificata": "data_pianificata",
    "data_pianificata": "data_pianificata",
    "riferimento esterno": "riferimento_esterno",
    "riferimento": "riferimento_esterno",
    "note": "note",
}


def _auto_map_col(col: str) -> str | None:
    col_clean = col.strip().lower()
    for template_col, field in FIELD_MAP.items():
        if col_clean == template_col.lower().replace(" *", ""):
            return field
    return _EXTRA_ALIASES.get(col_clean)


def _normalize_enum(value: str, valid: set[str], default: str) -> str:
    normalized = value.strip().upper().replace(" ", "_")
    return normalized if normalized in valid else default


def _parse_date(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.isoformat()
    raw = str(value).strip()
    if not raw:
        return None
    parsed = pd.to_datetime(raw, dayfirst=True, errors="coerce")
    if pd.isna(parsed):
        return None
    return parsed.to_pydatetime().isoformat()


async def _find_cliente(nome: str | None, codice: str | None) -> dict | None:
    query = {"attivo": True}
    if codice:
        doc = await db.clienti.find_one({**query, "codice_cliente": codice.strip()})
        if doc:
            return doc
    if nome:
        doc = await db.clienti.find_one({**query, "ragione_sociale": nome.strip()})
        if doc:
            return doc
        doc = await db.clienti.find_one({**query, "ragione_sociale": {"$regex": f"^{re.escape(nome.strip())}$", "$options": "i"}})
        if doc:
            return doc
    return None


def serialize(doc: dict) -> dict:
    d = dict(doc)
    d["id"] = str(d.pop("_id"))
    return d


def to_object_id(value: str) -> ObjectId:
    try:
        return ObjectId(value)
    except (bson_errors.InvalidId, Exception):
        raise HTTPException(status_code=404, detail="Intervento non trovato")


@router.get("/", response_model=List[InterventoResponse])
async def lista_interventi(
    stato: Optional[str] = None,
    origine: Optional[str] = None,
    attivo: bool = True,
    limit: int = Query(default=1000, ge=1, le=2000),
):
    filtro: dict = {"attivo": attivo}
    if stato:
        filtro["stato"] = stato
    if origine:
        filtro["origine"] = origine
    interventi = await db.interventi.find(filtro).sort("created_at", -1).to_list(limit)
    return [serialize(i) for i in interventi]


@router.get("/template")
def scarica_template_interventi():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Interventi"

    header_fill = PatternFill("solid", fgColor="1E3A5F")
    required_fill = PatternFill("solid", fgColor="2563EB")
    header_font = Font(color="FFFFFF", bold=True, size=11)

    for col_idx, (col_name, _, required) in enumerate(TEMPLATE_FIELDS, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = required_fill if required else header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[cell.column_letter].width = 24

    ws.row_dimensions[1].height = 35
    ws.cell(row=2, column=1, value="Manutenzione programmata caldaia")
    ws.cell(row=2, column=2, value="Controllo annuale impianto e verifica sicurezza")
    ws.cell(row=2, column=3, value="MANUTENZIONE")
    ws.cell(row=2, column=4, value="DA_PIANIFICARE")
    ws.cell(row=2, column=5, value="MEDIA")
    ws.cell(row=2, column=6, value="Rossi SpA")
    ws.cell(row=2, column=9, value="Via Roma 1")
    ws.cell(row=2, column=10, value="20100")
    ws.cell(row=2, column=11, value="Milano")
    ws.cell(row=2, column=12, value="MI")
    ws.cell(row=2, column=13, value="2026-06-06")
    ws.cell(row=2, column=14, value="2026-07-15")

    ws_note = wb.create_sheet("Istruzioni")
    ws_note["A1"] = "ISTRUZIONI"
    ws_note["A1"].font = Font(bold=True, size=14)
    ws_note["A3"] = "• Titolo e obbligatorio"
    ws_note["A4"] = "• Origine: TICKET | MANUTENZIONE | PREVENTIVO_ACCETTATO | MANUALE"
    ws_note["A5"] = "• Stato: APERTO | DA_PIANIFICARE | PIANIFICATO | IN_CORSO | COMPLETATO | ANNULLATO"
    ws_note["A6"] = "• Priorita: BASSA | MEDIA | ALTA | URGENTE"
    ws_note["A7"] = "• Scadenza: usa formato gg/mm/aaaa oppure aaaa-mm-gg"
    ws_note["A8"] = "• Cliente: se gia presente in anagrafica viene collegato automaticamente"

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=template_interventi.xlsx"}
    )


@router.post("/preview-import")
async def preview_import_interventi(file: UploadFile = File(...)):
    content = await file.read()
    try:
        if file.filename.endswith(".csv"):
            df = pd.read_csv(BytesIO(content), dtype=str, keep_default_na=False)
        else:
            df = pd.read_excel(BytesIO(content), dtype=str, keep_default_na=False)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Errore lettura file: {str(e)}")

    df = df.fillna("")
    auto_mapping: Dict[str, str] = {}
    for col in df.columns:
        field = _auto_map_col(col)
        if field:
            auto_mapping[col] = field

    return {
        "columns": list(df.columns),
        "auto_mapping": auto_mapping,
        "template_fields": [{"label": f[0], "field": f[1], "required": f[2]} for f in TEMPLATE_FIELDS],
        "preview": df.head(3).to_dict(orient="records"),
        "total_rows": len(df),
    }


@router.post("/import")
async def import_interventi(file: UploadFile = File(...), mapping: str = Form(default="")):
    import json

    content = await file.read()
    try:
        if file.filename.endswith(".csv"):
            df = pd.read_csv(BytesIO(content), dtype=str, keep_default_na=False)
        else:
            df = pd.read_excel(BytesIO(content), dtype=str, keep_default_na=False)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Errore lettura file: {str(e)}")

    df = df.fillna("")
    try:
        raw_map: Dict[str, str] = json.loads(mapping) if mapping else {}
        col_map = {k: v for k, v in raw_map.items() if v and v.strip()}
    except Exception:
        col_map = {}

    if not col_map:
        for col in df.columns:
            field = _auto_map_col(col)
            if field:
                col_map[col] = field

    valid_origini = {o.value for o in OrigineIntervento}
    valid_stati = {s.value for s in StatoIntervento}
    valid_priorita = {p.value for p in PrioritaIntervento}
    results: Dict[str, Any] = {"importati": 0, "saltati": 0, "errori": []}

    for idx, row in df.iterrows():
        try:
            data: Dict[str, Any] = {}
            for col, field in col_map.items():
                val = str(row.get(col, "")).strip()
                if val:
                    data[field] = val

            if not data.get("titolo"):
                results["saltati"] += 1
                continue

            cliente = await _find_cliente(data.get("cliente_nome"), data.get("cliente_codice"))
            sede = (cliente or {}).get("sede_legale") or {}
            cliente_nome = data.get("cliente_nome") or (cliente or {}).get("ragione_sociale")
            cliente_codice = data.get("cliente_codice") or (cliente or {}).get("codice_cliente")

            now = datetime.now(timezone.utc)
            doc = {
                "titolo": data.get("titolo"),
                "descrizione": data.get("descrizione") or None,
                "origine": _normalize_enum(data.get("origine", ""), valid_origini, OrigineIntervento.manutenzione.value),
                "stato": _normalize_enum(data.get("stato", ""), valid_stati, StatoIntervento.da_pianificare.value),
                "priorita": _normalize_enum(data.get("priorita", ""), valid_priorita, PrioritaIntervento.media.value),
                "cliente_id": str(cliente["_id"]) if cliente else None,
                "cliente_codice": cliente_codice,
                "cliente_nome": cliente_nome,
                "tecnico_id": None,
                "tecnico_nome": data.get("tecnico_nome") or None,
                "indirizzo": data.get("indirizzo") or sede.get("indirizzo"),
                "cap": data.get("cap") or sede.get("cap"),
                "citta": data.get("citta") or sede.get("citta"),
                "provincia": data.get("provincia") or sede.get("provincia"),
                "data_richiesta": _parse_date(data.get("data_richiesta")),
                "data_pianificata": _parse_date(data.get("data_pianificata")),
                "lat": None,
                "lng": None,
                "riferimento_esterno": data.get("riferimento_esterno") or None,
                "note": data.get("note") or None,
                "codice_intervento": await genera_codice_intervento(),
                "attivo": True,
                "created_at": now,
                "updated_at": now,
            }
            await db.interventi.insert_one(doc)
            results["importati"] += 1
        except Exception as e:
            results["errori"].append(f"Riga {idx + 2}: {str(e)}")

    return results


@router.delete("/clear-all")
async def svuota_interventi():
    result = await db.interventi.delete_many({})
    return {"ok": True, "eliminati": result.deleted_count}


@router.get("/{intervento_id}", response_model=InterventoResponse)
async def get_intervento(intervento_id: str):
    doc = await db.interventi.find_one({"_id": to_object_id(intervento_id)})
    if not doc:
        raise HTTPException(status_code=404, detail="Intervento non trovato")
    return serialize(doc)


@router.post("/", response_model=InterventoResponse, status_code=status.HTTP_201_CREATED)
async def crea_intervento(intervento: InterventoCreate):
    now = datetime.now(timezone.utc)
    data = intervento.model_dump(mode="json")
    data["codice_intervento"] = await genera_codice_intervento()
    data["attivo"] = True
    data["created_at"] = now
    data["updated_at"] = now
    result = await db.interventi.insert_one(data)
    doc = await db.interventi.find_one({"_id": result.inserted_id})
    return serialize(doc)


@router.put("/{intervento_id}", response_model=InterventoResponse)
async def aggiorna_intervento(intervento_id: str, intervento: InterventoUpdate):
    oid = to_object_id(intervento_id)
    doc = await db.interventi.find_one({"_id": oid})
    if not doc:
        raise HTTPException(status_code=404, detail="Intervento non trovato")
    aggiornamenti = {k: v for k, v in intervento.model_dump(mode="json").items() if v is not None}
    aggiornamenti["updated_at"] = datetime.now(timezone.utc)
    await db.interventi.update_one({"_id": oid}, {"$set": aggiornamenti})
    doc = await db.interventi.find_one({"_id": oid})
    return serialize(doc)


@router.patch("/{intervento_id}/stato")
async def aggiorna_stato(intervento_id: str, body: dict):
    stato = body.get("stato")
    if stato not in [s.value for s in StatoIntervento]:
        raise HTTPException(status_code=400, detail="Stato non valido")
    await db.interventi.update_one(
        {"_id": to_object_id(intervento_id)},
        {"$set": {"stato": stato, "updated_at": datetime.now(timezone.utc)}}
    )
    return {"ok": True}


@router.patch("/{intervento_id}/geo")
async def salva_geo(intervento_id: str, body: dict):
    lat = body.get("lat")
    lng = body.get("lng")
    if lat is None or lng is None:
        raise HTTPException(status_code=400, detail="lat e lng obbligatori")
    await db.interventi.update_one(
        {"_id": to_object_id(intervento_id)},
        {"$set": {"lat": lat, "lng": lng, "updated_at": datetime.now(timezone.utc)}}
    )
    return {"ok": True}


@router.delete("/{intervento_id}", status_code=status.HTTP_204_NO_CONTENT)
async def elimina_intervento(intervento_id: str):
    oid = to_object_id(intervento_id)
    doc = await db.interventi.find_one({"_id": oid})
    if not doc:
        raise HTTPException(status_code=404, detail="Intervento non trovato")
    await db.interventi.update_one(
        {"_id": oid},
        {"$set": {"attivo": False, "updated_at": datetime.now(timezone.utc)}}
    )
